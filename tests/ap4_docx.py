"""Shared helpers of the A-P4 tests (build day 10, lane A): the extra's skip rule, the
synthetic document, and DOCX text extraction through python-docx.

**The skip rule.** A test that needs docxtpl, python-docx or matplotlib carries
:data:`needs_extra`: it is skipped, with the reason below, when the extra is not installed
- unless ``PROOFPACK_REQUIRE_DOCX=1`` is set, in which case nothing is skipped and the
missing package fails the test. The CI job ``docx-extra`` (``.github/workflows/ci.yml``)
installs the extra and sets that variable, so the skip cannot hide there; the main pytest
job runs without the extra and the same tests skip with a named reason.

**Text extraction.** python-docx reads body paragraphs, table cells (each cell once, a
merged cell not repeated), and each section's header and footer. Customer text is
identified by style, the way the HTML identifies it by ``.customer-text``: paragraphs in
``PP Manufacturer Text`` / ``PP Manufacturer Text Label`` and runs in ``PP Manufacturer
Text Inline``. Status words are identified by runs in ``PP Status`` (the HTML's
``.status``).
"""

from __future__ import annotations

import copy
import io
import os
import re
import zipfile
from functools import cache
from pathlib import Path
from typing import Any

import pytest

from assembler import assemble
from conftest import make_criteria
from proofpack.render.docx import extra_available
from test_criteria import CRITERIA, FAIRNESS, cohort_with_a_thirty_row_site

REPO = Path(__file__).resolve().parent.parent
TEMPLATES = REPO / "src" / "proofpack" / "templates"
GENERATOR = REPO / "scripts" / "make_docx_templates.py"
REQUIRE_ENV = "PROOFPACK_REQUIRE_DOCX"
EXTRA_REQUIRED = os.environ.get(REQUIRE_ENV) == "1"
SKIP_REASON = (
    "the [docx] extra (docxtpl, python-docx, matplotlib) is not installed: pip install "
    '"proofpack[docx]"; the CI job docx-extra installs it and runs -m ap4 with '
    f"{REQUIRE_ENV}=1, where this skip is refused"
)
needs_extra = pytest.mark.skipif(not extra_available() and not EXTRA_REQUIRED, reason=SKIP_REASON)
NUMPY_MASK = "x.y.z"
MANUFACTURER_PARAGRAPH_STYLES = {"PP Manufacturer Text", "PP Manufacturer Text Label"}
MANUFACTURER_RUN_STYLE = "PP Manufacturer Text Inline"
STATUS_RUN_STYLE = "PP Status"
TIER_GLYPHS = ("ᵃ", "ᵇ", "ᶜ")


@cache
def _base_document() -> dict[str, Any]:
    crit = make_criteria(criteria=copy.deepcopy(CRITERIA), fairness=FAIRNESS)
    doc = assemble(cohort_with_a_thirty_row_site(), crit)
    doc["manifest"]["numpy"] = NUMPY_MASK
    return doc


def synthetic_document(
    *, watermark: str | None = None, data_marking: str | None = None
) -> dict[str, Any]:
    """The T1/T7/T8 golden tests' synthetic document (assembler, CRITERIA + FAIRNESS,
    B = 200, seed 20240101, ``manifest.numpy`` masked), a deep copy with the manifest's
    marks set as asked."""
    doc = copy.deepcopy(_base_document())
    doc["manifest"]["watermark"] = watermark
    if data_marking is not None:
        doc["manifest"]["data_marking"] = data_marking
    return doc


def load_generator():
    """``scripts/make_docx_templates.py`` as a module (it imports python-docx)."""
    import importlib.util
    import sys

    spec = importlib.util.spec_from_file_location("make_docx_templates_ap4", GENERATOR)
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    sys.modules["make_docx_templates_ap4"] = module
    spec.loader.exec_module(module)
    return module


# ------------------------------------------------------------------ the package


def part(data: bytes, name: str) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return z.read(name).decode("utf-8")


def part_names(data: bytes) -> list[str]:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return z.namelist()


def footer_parts(data: bytes) -> list[str]:
    names = sorted(
        (n for n in part_names(data) if re.fullmatch(r"word/footer\d+\.xml", n)),
        key=lambda n: int(re.search(r"\d+", n).group(0)),
    )
    return [part(data, n) for n in names]


def header_parts(data: bytes) -> list[str]:
    names = sorted(
        (n for n in part_names(data) if re.fullmatch(r"word/header\d+\.xml", n)),
        key=lambda n: int(re.search(r"\d+", n).group(0)),
    )
    return [part(data, n) for n in names]


def xml_text(xml: str) -> str:
    """The text of one XML part: ``<w:t>`` runs joined, paragraphs separated by ``\\n``."""
    out = []
    for para in re.findall(r"<w:p\b.*?</w:p>", xml, re.S):
        out.append("".join(re.findall(r"<w:t(?:\s[^>]*)?>([^<]*)</w:t>", para)))
    return "\n".join(_unescape(t) for t in out)


def _unescape(s: str) -> str:
    return (
        s.replace("&lt;", "<")
        .replace("&gt;", ">")
        .replace("&quot;", '"')
        .replace("&apos;", "'")
        .replace("&amp;", "&")
    )


# ------------------------------------------------------------------ python-docx reads


def open_docx(data: bytes | Path):
    import docx  # noqa: PLC0415 - the extra

    return docx.Document(io.BytesIO(data) if isinstance(data, bytes) else str(data))


def unique_cells(table) -> list[Any]:
    """Each cell of ``table`` once (a merged cell repeats in ``row.cells``)."""
    seen: set[int] = set()
    out = []
    for row in table.rows:
        for cell in row.cells:
            key = id(cell._tc)
            if key in seen:
                continue
            seen.add(key)
            out.append(cell)
    return out


def body_paragraphs(d) -> list[Any]:
    """Every paragraph of the body: the top-level ones and those inside table cells."""
    out = list(d.paragraphs)
    for table in d.tables:
        for cell in unique_cells(table):
            out.extend(cell.paragraphs)
    return out


def paragraph_texts(data: bytes | Path) -> list[str]:
    return [p.text for p in body_paragraphs(open_docx(data))]


def cell_texts(data: bytes | Path) -> list[str]:
    d = open_docx(data)
    return [c.text for t in d.tables for c in unique_cells(t)]


def section_footers(data: bytes | Path) -> list[str]:
    d = open_docx(data)
    return ["\n".join(p.text for p in s.footer.paragraphs) for s in d.sections]


def section_headers(data: bytes | Path) -> list[str]:
    d = open_docx(data)
    return ["\n".join(p.text for p in s.header.paragraphs) for s in d.sections]


def styled_runs(data: bytes | Path) -> list[tuple[str, str, str | None]]:
    """``(text, paragraph style, run style)`` for every run of every body paragraph
    (tables included)."""
    out = []
    for p in body_paragraphs(open_docx(data)):
        pstyle = p.style.name if p.style is not None else ""
        for r in p.runs:
            rstyle = r.style.name if r.style is not None else None
            if rstyle == "Default Paragraph Font":
                rstyle = None
            out.append((r.text, pstyle, rstyle))
    return out


def engine_text(data: bytes | Path) -> str:
    """The body text that is the engine's: every run outside the manufacturer styles."""
    return "\n".join(
        text
        for text, pstyle, rstyle in styled_runs(data)
        if pstyle not in MANUFACTURER_PARAGRAPH_STYLES and rstyle != MANUFACTURER_RUN_STYLE
    )


def all_text(data: bytes | Path) -> str:
    """Body paragraphs and cells, headers and footers, joined."""
    return "\n".join(paragraph_texts(data) + section_headers(data) + section_footers(data))


def words(text: str) -> set[str]:
    return set(re.split(r"[^a-z]+", text.lower())) - {""}
